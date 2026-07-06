#!/usr/bin/env python3
"""Generate a single portrait .docx with all Scenario A result tables
(Latency, Throughput, CPU, RAM) for every payload already collected in
results/scenario_a/raw.csv.

Table look matches the thesis's existing Tabel 3.2 style: light gray
header, alternating (zebra) row bands, bold first column. Ratio cells
(P50/P95/P99) get the same goodness-based color highlight used in the
HTML tables, overriding the zebra band for that cell only.

Requires python-docx (not in the standard library):
    pip install python-docx
"""
import argparse
import csv
import math
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from collect_scenario_a import (
    CANONICAL_SEGMENTS,
    LOWER_IS_BETTER,
    METRIC_DISPLAY_NAME,
    METRIC_SHORT_NAME,
    METRICS,
    fmt,
)

PAYLOADS = ["10mb", "50mb", "100mb"]

HEADER_BG = "D9D9D9"
BAND_BG = "F2F2F2"
RATIO_COLORS = {
    "neutral": ("EEF0F2", "000000"),
    "good-light": ("DCEFE4", "000000"),
    "good-medium": ("A9D9BD", "000000"),
    "good-strong": ("3F8F6C", "FFFFFF"),
    "warn-light": ("F7E3DA", "000000"),
    "warn-medium": ("EAB596", "000000"),
    "warn-strong": ("B5502F", "FFFFFF"),
    "na": ("EDEDED", "595959"),
}

CPU_RAM_NOTES = [
    'Baris "Worker ke Gateway" tidak punya data CPU/RAM.',
    "Arsitektur sistem sengaja tidak mengukur CPU/RAM pada segmen ini.",
    "Baris tetap ditampilkan supaya bentuk tabel konsisten dengan tabel lain.",
]


def ratio_band(ratio, lower_is_better):
    if ratio is None or ratio <= 0:
        return RATIO_COLORS["na"]
    log = math.log2(ratio)
    if not lower_is_better:
        log = -log
    if abs(log) < 0.15:
        return RATIO_COLORS["neutral"]
    direction = "warn" if log > 0 else "good"
    mag = abs(log)
    tier = "light" if mag < 1 else "medium" if mag < math.log2(5) else "strong"
    return RATIO_COLORS[f"{direction}-{tier}"]


def load_raw_by_payload(csv_path):
    raw_by_payload = defaultdict(dict)
    with open(csv_path, newline="") as f:
        for row in csv.DictReader(f):
            key = (row["metric"], row["segment"], row["protocol"])
            pct = int(row["percentile"])
            raw_by_payload[row["payload_size"]].setdefault(key, {})[pct] = float(row["value"])
    return raw_by_payload


def build_rows(raw, metric):
    rows = []
    for segment_key, segment_label in CANONICAL_SEGMENTS:
        rest = raw.get((metric, segment_key, "rest"), {})
        grpc = raw.get((metric, segment_key, "grpc"), {})
        row = [segment_label]
        for pct in (50, 95, 99):
            row.append(rest.get(pct))
        for pct in (50, 95, 99):
            row.append(grpc.get(pct))
        for pct in (50, 95, 99):
            r, g = rest.get(pct), grpc.get(pct)
            row.append(round(g / r, 2) if (r and g) else None)
        rows.append(row)
    return rows


def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def style_cell(cell, *, bold=False, size=8, align=WD_ALIGN_PARAGRAPH.CENTER, color=None):
    for p in cell.paragraphs:
        p.alignment = align
        for run in p.runs:
            run.font.size = Pt(size)
            run.bold = bold
            if color:
                run.font.color.rgb = RGBColor.from_string(color)


COL_WIDTHS = [Cm(2.4)] + [Cm(1.25)] * 9  # ~13.65cm total, fits a narrow portrait margin


def add_metric_table(doc, display_name, unit, rows, lower_is_better, extra_notes=None):
    doc.add_heading(f"{display_name} ({unit})", level=4)

    table = doc.add_table(rows=2 + len(rows), cols=10)
    table.style = doc.styles["Table Grid"]
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    seg_cell = table.cell(0, 0).merge(table.cell(1, 0))
    seg_cell.text = "Segmen"
    rest_cell = table.cell(0, 1).merge(table.cell(0, 3))
    rest_cell.text = "REST"
    grpc_cell = table.cell(0, 4).merge(table.cell(0, 6))
    grpc_cell.text = "gRPC"
    ratio_cell = table.cell(0, 7).merge(table.cell(0, 9))
    ratio_cell.text = "Rasio (gRPC/REST)"

    for i, label in enumerate(["P50", "P95", "P99"] * 3, start=1):
        table.cell(1, i).text = label

    for r in range(2):
        for c in range(10):
            set_cell_shading(table.cell(r, c), HEADER_BG)
            style_cell(table.cell(r, c), bold=True)

    for ridx, row in enumerate(rows):
        r = 2 + ridx
        banded = ridx % 2 == 1
        table.cell(r, 0).text = row[0]
        style_cell(table.cell(r, 0), bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
        if banded:
            set_cell_shading(table.cell(r, 0), BAND_BG)

        for c in range(1, 7):
            cell = table.cell(r, c)
            cell.text = fmt(row[c])
            style_cell(cell)
            if banded:
                set_cell_shading(cell, BAND_BG)

        for c in range(7, 10):
            cell = table.cell(r, c)
            cell.text = fmt(row[c], "×")
            bg, fg = ratio_band(row[c], lower_is_better)
            set_cell_shading(cell, bg)
            style_cell(cell, color=fg)

    for r in range(2 + len(rows)):
        for c, w in enumerate(COL_WIDTHS):
            table.cell(r, c).width = w

    notes = [
        "Rasio dihitung sebagai nilai gRPC dibagi nilai REST pada persentil yang sama.",
        "Nilai lebih rendah lebih baik untuk metrik ini."
        if lower_is_better
        else "Nilai lebih tinggi lebih baik untuk metrik ini.",
    ] + (extra_notes or [])
    for note in notes:
        p = doc.add_paragraph(note, style="List Bullet")
        style_cell_text_size(p, 8)

    doc.add_paragraph()


def style_cell_text_size(paragraph, size):
    for run in paragraph.runs:
        run.font.size = Pt(size)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--root",
        default=str(Path(__file__).resolve().parent.parent),
        help="repo root, results/scenario_a/raw.csv is read from here",
    )
    parser.add_argument(
        "--output",
        default=None,
        help="output .docx path, defaults to results/scenario_a/tabel_hasil_skenario_a.docx",
    )
    args = parser.parse_args()

    root = Path(args.root)
    csv_path = root / "results" / "scenario_a" / "raw.csv"
    raw_by_payload = load_raw_by_payload(csv_path)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.size = Pt(9)

    doc.add_heading("Tabel Hasil Skenario A", level=1)

    for payload in PAYLOADS:
        raw = raw_by_payload.get(payload)
        if not raw:
            continue
        doc.add_heading(f"Payload {payload.upper()}", level=2)
        for metric, unit, _factor in METRICS:
            short = METRIC_SHORT_NAME[metric]
            display = METRIC_DISPLAY_NAME[short]
            rows = build_rows(raw, metric)
            extra = CPU_RAM_NOTES if short in ("cpu", "ram") else None
            add_metric_table(doc, display, unit, rows, LOWER_IS_BETTER[short], extra)

    output = Path(args.output) if args.output else root / "results" / "scenario_a" / "tabel_hasil_skenario_a.docx"
    doc.save(output)
    print(f"Wrote {output}")


if __name__ == "__main__":
    main()
